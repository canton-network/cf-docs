"""Interactive report tabs for the cf-docs contribution tracker (Google Sheets).

Each tab is driven by dropdowns at the top. The figures are Google Sheets formulas
(LET, FILTER, SORT, HSTACK, ARRAYFORMULA, SWITCH) over the data tabs, so they recompute in
Sheets whenever a dropdown, a date or the Contributors affiliations change. They are
written into the .xlsx that is imported into Sheets; LibreOffice cannot evaluate all of
them, so the update script checks these tabs on Google's side instead.

Ranges are fixed to the current data size (the .xlsx format has no open-ended ranges);
the tabs are regenerated with each data refresh, so they always cover every row.
"""

from __future__ import annotations

from dataclasses import dataclass

from openpyxl.chart import BarChart, LineChart, Reference, Series
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
INK, INK2, INK3 = "141A24", "4B5566", "7D8798"
ACCENT = "1C5CAB"
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
TILE_FILL = PatternFill("solid", fgColor="F3F5F8")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
LINE = Side(style="thin", color="D9DDE5")
HAIR = Side(style="hair", color="E8EBF0")

PERIODS = ["Last week", "Last 4 weeks", "Last 13 weeks", "This month", "Last month", "This quarter",
           "Last quarter", "Year to date", "All time", "Custom"]
AFFILIATIONS = ["Everyone", "Digital Asset", "Canton Foundation", "External"]
MAX_BUCKETS = 156  # three years of weeks


@dataclass
class Ranges:
    p_last: int
    r_last: int
    i_last: int
    c_last: int
    asof: str = "_meta!$B$2"

    def P(self, col: str) -> str:
        return f"PRs!${col}$2:${col}${self.p_last}"

    def R(self, col: str) -> str:
        return f"Reviews!${col}$2:${col}${self.r_last}"

    def I(self, col: str) -> str:  # noqa: E743
        return f"Issues!${col}$2:${col}${self.i_last}"

    def C(self, col: str) -> str:
        return f"Contributors!${col}$2:${col}${self.c_last}"


def font(size: float = 10, bold: bool = False, color: str = INK, italic: bool = False) -> Font:
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


# ---------- Shared layout pieces ----------

def masthead(ws, title_formula: str, subtitle_formula: str) -> None:
    for col in "ABCDEFGHIJ":
        ws.column_dimensions[col].width = 14
    ws.column_dimensions["B"].width = 18
    ws["A1"] = "CANTON-NETWORK / CF-DOCS · CONTRIBUTION TRACKER"
    ws["A1"].font = font(9, bold=True, color=INK3)
    ws.merge_cells("A2:J2")
    ws["A2"] = title_formula
    ws["A2"].font = font(20, bold=True)
    ws.row_dimensions[2].height = 32
    ws.merge_cells("A3:J3")
    ws["A3"] = subtitle_formula
    ws["A3"].font = font(9, color=INK2)


def dropdown(ws, label_ref: str, cell_ref: str, label: str, options: list[str], default: str) -> None:
    ws[label_ref] = label
    ws[label_ref].font = font(8, bold=True, color=INK3)
    ws[cell_ref] = default
    ws[cell_ref].font = font(10, bold=True, color="0000FF")
    ws[cell_ref].fill = INPUT_FILL
    ws[cell_ref].border = Border(left=LINE, right=LINE, top=LINE, bottom=LINE)
    validation = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=False,
                                showDropDown=False)
    ws.add_data_validation(validation)
    validation.add(cell_ref)


def date_input(ws, label_ref: str, cell_ref: str, label: str, value: str) -> None:
    ws[label_ref] = label
    ws[label_ref].font = font(8, bold=True, color=INK3)
    ws[cell_ref] = value
    ws[cell_ref].number_format = "d mmm yyyy"
    ws[cell_ref].font = font(10, color="0000FF")
    ws[cell_ref].fill = INPUT_FILL
    ws[cell_ref].border = Border(left=LINE, right=LINE, top=LINE, bottom=LINE)
    validation = DataValidation(type="date", operator="greaterThan", formula1="DATE(2020,1,1)")
    ws.add_data_validation(validation)
    validation.add(cell_ref)


def period_cells(ws, r: Ranges, period: str, custom_from: str, custom_to: str, helper: str = "Z") -> dict:
    """Hidden helper cells holding the chosen period; returns their absolute references."""
    a = r.asof
    month0 = f"DATE(YEAR({a}),MONTH({a}),1)"
    quarter0 = f"DATE(YEAR({a}),3*INT((MONTH({a})-1)/3)+1,1)"
    h = helper
    ws[f"{h}1"] = "helper"
    # Most recent Monday 12:00 at or before the refresh time: weeks run Monday noon to Monday noon.
    ws[f"{h}2"] = f"=LET(d,INT({a})-WEEKDAY(INT({a}),3)+0.5,IF(d>{a},d-7,d))"
    ws[f"{h}3"] = (f'=SWITCH({period},"Last week",{h}2-7,"Last 4 weeks",{h}2-28,"Last 13 weeks",{h}2-91,'
                   f'"This month",{month0},"Last month",EDATE({month0},-1),"This quarter",{quarter0},'
                   f'"Last quarter",EDATE({quarter0},-3),"Year to date",DATE(YEAR({a}),1,1),'
                   f'"All time",INT(MIN({r.P("F")},{r.I("E")})),"Custom",{custom_from},{h}2-7)')
    ws[f"{h}4"] = (f'=SWITCH({period},"Last week",{h}2,"Last 4 weeks",{h}2,"Last 13 weeks",{h}2,'
                   f'"Last month",{month0},"Last quarter",{quarter0},"Custom",{custom_to}+1,{a})')
    ws[f"{h}5"] = f"={h}3-({h}4-{h}3)"
    for k in range(2, 6):
        ws[f"{h}{k}"].number_format = "yyyy-mm-dd hh:mm"
    ws.column_dimensions[h].hidden = True
    return {"s": f"${h}$3", "e": f"${h}$4", "ps": f"${h}$5"}


def period_title(p: dict) -> str:
    return f'=TEXT({p["s"]},"d mmm yyyy")&" – "&TEXT({p["e"]}-1/86400,"d mmm yyyy")'


def period_subtitle(p: dict, period: str, r: Ranges, extra: str = '""') -> str:
    return (f'={period}&" · compared with "&TEXT({p["ps"]},"d mmm")&" – "&TEXT({p["s"]}-1/86400,"d mmm")'
            f'&" · times US Eastern · data refreshed "&TEXT({r.asof},"d mmm yyyy h:mm AM/PM")&" ET"&{extra}')


def tile(ws, row: int, col: int, value: str, label: str, sub: str, fmt: str = "0") -> None:
    from openpyxl.utils import get_column_letter
    c1, c2 = get_column_letter(col), get_column_letter(col + 1)
    for k in range(3):
        ws.merge_cells(f"{c1}{row + k}:{c2}{row + k}")
        for cc in (col, col + 1):
            cell = ws.cell(row=row + k, column=cc)
            cell.fill = TILE_FILL
            cell.border = Border(left=LINE if cc == col else None, right=LINE if cc == col + 1 else None,
                                 top=LINE if k == 0 else None, bottom=LINE if k == 2 else None)
    ws[f"{c1}{row}"] = value
    ws[f"{c1}{row}"].number_format = fmt
    ws[f"{c1}{row}"].font = font(20, bold=True)
    ws[f"{c1}{row}"].alignment = Alignment(horizontal="left", indent=1)
    ws[f"{c1}{row + 1}"] = label
    ws[f"{c1}{row + 1}"].font = font(9, color=INK2)
    ws[f"{c1}{row + 1}"].alignment = Alignment(indent=1)
    ws[f"{c1}{row + 2}"] = sub
    ws[f"{c1}{row + 2}"].font = font(9, color=INK3)
    ws[f"{c1}{row + 2}"].alignment = Alignment(indent=1)
    ws.row_dimensions[row].height = 30


def compare(cur: str, old: str) -> str:
    return (f'=IFERROR(IF(({cur})>({old}),"▲ from "&({old}),IF(({cur})<({old}),"▼ from "&({old}),'
            f'"same as previous period")),"")')


def section(ws, row: int, title: str, note: str = "") -> None:
    ws.cell(row=row, column=1, value=title).font = font(13, bold=True)
    if note:
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=10)
        ws.cell(row=row, column=3, value=note).font = font(9, color=INK3)
    ws.row_dimensions[row].height = 22


def table_header(ws, row: int, labels: list[str], first_col: int = 1) -> None:
    for j, label in enumerate(labels, start=first_col):
        cell = ws.cell(row=row, column=j, value=label)
        cell.font = font(8, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 28


def body_rows(ws, first: int, count: int, cols: int, formats: dict[int, str] | None = None) -> None:
    for i in range(first, first + count):
        for j in range(1, cols + 1):
            cell = ws.cell(row=i, column=j)
            cell.font = font(10, color=ACCENT if j == 1 else INK)
            cell.border = Border(bottom=HAIR)
            if formats and j in formats:
                cell.number_format = formats[j]


def notes(ws, row: int, lines: list[str]) -> None:
    for k, text in enumerate(lines):
        ws.merge_cells(start_row=row + k, start_column=1, end_row=row + k, end_column=10)
        ws.cell(row=row + k, column=1, value=text).font = font(8, color=INK3)


def aff_ok(rng: str, choice: str) -> str:
    """Array condition: the affiliation matches the dropdown, or everyone except bots."""
    return f'((({rng}={choice})+({choice}="Everyone")*({rng}<>"Bot"))>0)'


def aff_crit(choice: str) -> str:
    """COUNTIFS criterion for the same choice."""
    return f'IF({choice}="Everyone","<>Bot",{choice})'


# ---------- Activity ----------

def activity(ws, r: Ranges) -> None:
    dropdown(ws, "A5", "A6", "PERIOD", PERIODS, "Last 13 weeks")
    dropdown(ws, "C5", "C6", "GROUP TREND BY", ["Week", "Month", "Quarter"], "Week")
    dropdown(ws, "E5", "E6", "AFFILIATION", AFFILIATIONS, "Everyone")
    date_input(ws, "G5", "G6", "CUSTOM FROM", f"=EDATE(DATE(YEAR({r.asof}),MONTH({r.asof}),1),-1)")
    date_input(ws, "I5", "I6", "CUSTOM TO (INCLUSIVE)", f"=DATE(YEAR({r.asof}),MONTH({r.asof}),1)-1")
    ws.merge_cells("A6:B6"), ws.merge_cells("C6:D6"), ws.merge_cells("E6:F6"), ws.merge_cells("G6:H6"), \
        ws.merge_cells("I6:J6")
    p = period_cells(ws, r, "$A$6", "$G$6", "$I$6")
    masthead(ws, period_title(p), period_subtitle(p, "$A$6", r,
             f'IF($E$6="Everyone",""," · "&$E$6&" only")'))
    s, e, ps, aff = p["s"], p["e"], p["ps"], "$E$6"
    crit = aff_crit(aff)

    def between(rng: str, a: str, b: str) -> str:
        return f'{rng},">="&{a},{rng},"<"&{b}'

    m = {
        "opened": lambda a, b: f'COUNTIFS({between(r.P("F"), a, b)},{r.P("D")},{crit})',
        "merged": lambda a, b: f'COUNTIFS({between(r.P("G"), a, b)},{r.P("D")},{crit})',
        "closed": lambda a, b: f'COUNTIFS({between(r.P("H"), a, b)},{r.P("G")},"",{r.P("D")},{crit})',
        "first_review": lambda a, b: (f'IFERROR(MEDIAN(FILTER({r.P("M")},{r.P("F")}>={a},{r.P("F")}<{b},'
                                      f'{r.P("M")}<>"",{aff_ok(r.P("D"), aff)})),"–")'),
        "reviewed": lambda a, b: f'COUNTIFS({between(r.R("E"), a, b)},{r.R("C")},{crit},{r.R("F")},TRUE)',
        "issues_opened": lambda a, b: f'COUNTIFS({between(r.I("E"), a, b)},{r.I("D")},{crit})',
        "issues_closed": lambda a, b: f'COUNTIFS({between(r.I("F"), a, b)},{r.I("D")},{crit})',
        "open_issues": lambda a, b: (f'COUNTIFS({r.I("E")},"<"&{b},{r.I("F")},"",{r.I("D")},{crit})'
                                     f'+COUNTIFS({r.I("E")},"<"&{b},{r.I("F")},">="&{b},{r.I("D")},{crit})'),
    }
    tiles = [
        ("opened", "PRs opened"), ("merged", "PRs merged"), ("closed", "PRs closed without merging"),
        ("first_review", "Median hours to first review"), ("reviewed", "PRs reviewed (reviewer–PR pairs)"),
        ("issues_opened", "Issues opened"), ("issues_closed", "Issues closed"),
        ("open_issues", "Open issues at period end"),
    ]
    auto_closed = {
        "closed": f'COUNTIFS({between(r.P("H"), s, e)},{r.P("G")},"",{r.P("D")},{crit},{r.P("S")},TRUE)',
        "issues_closed": f'COUNTIFS({between(r.I("F"), s, e)},{r.I("D")},{crit},{r.I("M")},TRUE)',
    }
    for k, (key, label) in enumerate(tiles):
        cur, old = m[key](s, e), m[key](ps, s)
        sub = (f'=IFERROR("previous period "&TEXT({old},"0.0"),"")' if key == "first_review"
               else compare(cur, old))
        if key in auto_closed:
            label = f'="{label}"&IF({auto_closed[key]}>0," · "&{auto_closed[key]}&" by the follow-up automation","")'
        tile(ws, 8 + 4 * (k // 4), 1 + 2 * (k % 4), f"={cur}", label, sub,
             "0.0" if key == "first_review" else "0")

    row = 37
    section(ws, row, "Trend table", "One row per week, month or quarter; the first and last may be partial.")
    head = ["Period", "PRs opened", "PRs merged", "PRs closed without merging", "PRs reviewed", "Issues opened",
            "Issues closed"]
    table_header(ws, row + 1, head)
    first = row + 2
    last_day = f"({e}-1/86400)"
    n = (f'MIN({MAX_BUCKETS},IF($C$6="Week",ROUNDUP(({e}-{s})/7,0),IF($C$6="Month",'
         f'(YEAR({last_day})-YEAR({s}))*12+MONTH({last_day})-MONTH({s})+1,'
         f'(YEAR({last_day})-YEAR({s}))*4+INT((MONTH({last_day})-1)/3)-INT((MONTH({s})-1)/3)+1)))')
    month1 = f"DATE(YEAR({s}),MONTH({s}),1)"
    quarter1 = f"DATE(YEAR({s}),3*INT((MONTH({s})-1)/3)+1,1)"
    bound = lambda k: (f'IF($C$6="Week",{s}+7*({k}),IF($C$6="Month",EDATE({month1},{k}),'  # noqa: E731
                       f'EDATE({quarter1},3*({k}))))')
    pc = lambda col: r.P(col)  # noqa: E731
    ws.cell(row=first, column=1, value=(
        f'=ARRAYFORMULA(LET(k,SEQUENCE({n})-1,'
        f'st,IF({bound("k")}<{s},{s},{bound("k")}),'
        f'en,IF({bound("k+1")}>{e},{e},{bound("k+1")}),'
        # Label each bucket by the days it covers: "14–21 Sep", or "Sep 2026" / "Q3 2026" when whole.
        # (LET names must not match function names: a variable called "month" hides MONTH().)
        f'lastt,en-1/86400,'
        f'daysx,IF(MONTH(st)=MONTH(lastt),TEXT(st,"d")&"–"&TEXT(lastt,"d mmm"),'
        f'TEXT(st,"d mmm")&"–"&TEXT(lastt,"d mmm")),'
        f'wholem,($C$6="Month")*(DAY(st)=1)*(st=INT(st))*(en=EDATE(st,1)),'
        f'wholeq,($C$6="Quarter")*(DAY(st)=1)*(st=INT(st))*(MOD(MONTH(st)-1,3)=0)*(en=EDATE(st,3)),'
        f'labels,IF(wholem,TEXT(st,"mmm yyyy"),IF(wholeq,"Q"&(INT((MONTH(st)-1)/3)+1)&" "&YEAR(st),daysx)),'
        f'HSTACK(labels,'
        f'COUNTIFS({pc("F")},">="&st,{pc("F")},"<"&en,{pc("D")},{crit}),'
        f'COUNTIFS({pc("G")},">="&st,{pc("G")},"<"&en,{pc("D")},{crit}),'
        f'COUNTIFS({pc("H")},">="&st,{pc("H")},"<"&en,{pc("G")},"",{pc("D")},{crit}),'
        f'COUNTIFS({r.R("E")},">="&st,{r.R("E")},"<"&en,{r.R("C")},{crit},{r.R("F")},TRUE),'
        f'COUNTIFS({r.I("E")},">="&st,{r.I("E")},"<"&en,{r.I("D")},{crit}),'
        f'COUNTIFS({r.I("F")},">="&st,{r.I("F")},"<"&en,{r.I("D")},{crit}))))'))

    chart = BarChart()
    chart.type, chart.grouping = "col", "clustered"
    chart.title = None  # the heading above the chart follows the dropdowns; a chart title cannot
    chart.y_axis.title = "Count"
    chart.legend.position = "b"
    chart.height, chart.width = 8, 25
    # PR measures in blues, issue measures in oranges.
    colors = ["1F4E8C", "4A86C8", "A9C6E8", "5B6C8F", "E07B39", "F2B880"]
    for col, color in zip(range(2, 8), colors):
        series = Series(Reference(ws, min_col=col, min_row=first, max_row=first + MAX_BUCKETS - 1),
                        title=head[col - 1])
        series.graphicalProperties.solidFill = color
        series.graphicalProperties.line.solidFill = color
        chart.series.append(series)
    chart.set_categories(Reference(ws, min_col=1, min_row=first, max_row=first + MAX_BUCKETS - 1))
    group = 'LOWER($C$6)'
    section(ws, 16, (f'="PRs and issues per "&{group}&", "&TEXT({s},"d mmm")&" – "&TEXT({e}-1/86400,"d mmm yyyy")'
                     f'&IF({aff}="Everyone",", everyone except bots",", "&{aff}&" only")'))
    ws.add_chart(chart, "A17")
    ws.merge_cells("A36:J36")
    ws["A36"] = (f'=IF({n.replace("MIN(" + str(MAX_BUCKETS) + ",", "(")}>{MAX_BUCKETS},'
                 f'"The trend shows only the first {MAX_BUCKETS} rows; choose a coarser grouping to see the whole period.","")')
    ws["A36"].font = font(9, bold=True, color="B3261E")
    notes(ws, 33, [
        "Affiliation filters by the PR or issue author (and by reviewer for PRs reviewed). Everyone means everyone except bots.",
        "Bars are labelled with the days they cover. Weeks run Monday 12:00 to Monday 12:00 US Eastern; months and quarters are calendar periods.",
        "Change the yellow cells. Custom uses the two dates on the right; the end date is included.",
    ])
    ws.freeze_panes = "A7"


# ---------- People ----------

def people(ws, r: Ranges) -> None:
    dropdown(ws, "A5", "A6", "PERIOD", PERIODS, "Last 4 weeks")
    dropdown(ws, "C5", "C6", "AFFILIATION", AFFILIATIONS, "Everyone")
    dropdown(ws, "E5", "E6", "SORT BY", ["PRs opened", "PRs merged", "PRs reviewed", "Issues opened"], "PRs opened")
    date_input(ws, "G5", "G6", "CUSTOM FROM", f"=EDATE(DATE(YEAR({r.asof}),MONTH({r.asof}),1),-1)")
    date_input(ws, "I5", "I6", "CUSTOM TO (INCLUSIVE)", f"=DATE(YEAR({r.asof}),MONTH({r.asof}),1)-1")
    for span in ("A6:B6", "C6:D6", "E6:F6", "G6:H6", "I6:J6"):
        ws.merge_cells(span)
    p = period_cells(ws, r, "$A$6", "$G$6", "$I$6")
    masthead(ws, period_title(p), period_subtitle(p, "$A$6", r))
    s, e, aff = p["s"], p["e"], "$C$6"

    row = 8
    section(ws, row, "By affiliation", "Bots excluded.")
    table_header(ws, row + 1, ["Affiliation", "People active", "PRs opened", "PRs merged", "PRs reviewed",
                               "Issues opened"])
    for k, name in enumerate(["Digital Asset", "Canton Foundation", "External"]):
        i = row + 2 + k
        q = f'"{name}"'
        ws.cell(row=i, column=1, value=name)
        # An empty FILTER is an error, not zero rows: blank those out before counting.
        ws.cell(row=i, column=2, value=(
            f'=IFERROR(ROWS(UNIQUE(LET(x,VSTACK('
            f'IFERROR(FILTER({r.P("C")},{r.P("D")}={q},{r.P("F")}>={s},{r.P("F")}<{e}),""),'
            f'IFERROR(FILTER({r.R("B")},{r.R("C")}={q},{r.R("E")}>={s},{r.R("E")}<{e}),""),'
            f'IFERROR(FILTER({r.I("C")},{r.I("D")}={q},{r.I("E")}>={s},{r.I("E")}<{e}),"")),'
            f'FILTER(x,x<>"")))),0)'))
        ws.cell(row=i, column=3, value=f'=COUNTIFS({r.P("D")},{q},{r.P("F")},">="&{s},{r.P("F")},"<"&{e})')
        ws.cell(row=i, column=4, value=f'=COUNTIFS({r.P("D")},{q},{r.P("G")},">="&{s},{r.P("G")},"<"&{e})')
        ws.cell(row=i, column=5, value=(f'=COUNTIFS({r.R("C")},{q},{r.R("E")},">="&{s},{r.R("E")},"<"&{e},'
                                        f'{r.R("F")},TRUE)'))
        ws.cell(row=i, column=6, value=f'=COUNTIFS({r.I("D")},{q},{r.I("E")},">="&{s},{r.I("E")},"<"&{e})')
    body_rows(ws, row + 2, 3, 6)
    for i in range(row + 2, row + 5):
        ws.cell(row=i, column=1).font = font(10)

    row = 14
    section(ws, row, "People", "Everyone with activity in the period. Change affiliations on the Contributors tab.")
    table_header(ws, row + 1, ["Login", "Affiliation", "PRs opened", "PRs merged", "PRs reviewed",
                               "Issues opened"])
    logins = f'FILTER({r.C("A")},{r.C("B")}<>"Bot",{aff_ok(r.C("B"), aff)})'
    ws.cell(row=row + 2, column=1, value=(
        f'=IFERROR(ARRAYFORMULA(LET(who,{logins},'
        f'af,VLOOKUP(who,Contributors!$A$2:$B${r.c_last},2,FALSE),'
        f'op,COUNTIFS({r.P("C")},who,{r.P("F")},">="&{s},{r.P("F")},"<"&{e}),'
        f'mg,COUNTIFS({r.P("C")},who,{r.P("G")},">="&{s},{r.P("G")},"<"&{e}),'
        f'rv,COUNTIFS({r.R("B")},who,{r.R("E")},">="&{s},{r.R("E")},"<"&{e},{r.R("F")},TRUE),'
        f'io,COUNTIFS({r.I("C")},who,{r.I("E")},">="&{s},{r.I("E")},"<"&{e}),'
        f'SORT(FILTER(HSTACK(who,af,op,mg,rv,io),op+mg+rv+io>0),'
        f'MATCH($E$6,{{"PRs opened","PRs merged","PRs reviewed","Issues opened"}},0)+2,FALSE))),'
        f'"No activity in this period.")'))
    ws.column_dimensions["A"].width = 24
    ws.freeze_panes = "A7"


# ---------- Open work ----------

def open_work(ws, r: Ranges) -> None:
    dropdown(ws, "A5", "A6", "SHOW", ["Pull requests", "Issues"], "Pull requests")
    dropdown(ws, "C5", "C6", "AFFILIATION", AFFILIATIONS, "Everyone")
    dropdown(ws, "E5", "E6", "UNTOUCHED FOR AT LEAST", ["Any time", "7 days", "30 days", "90 days"], "Any time")
    dropdown(ws, "G5", "G6", "DRAFT PRS", ["Exclude drafts", "Include drafts"], "Exclude drafts")
    dropdown(ws, "I5", "I6", "SORT", ["Least recently touched", "Newest first"], "Least recently touched")
    for span in ("A6:B6", "C6:D6", "E6:F6", "G6:H6", "I6:J6"):
        ws.merge_cells(span)
    a = r.asof
    masthead(ws, '="Open work right now"',
             f'="Not tied to a period. Untouched = days since a person last commented, committed, reviewed, labelled or '
             f'otherwise changed it; bots, including the follow-up automation, are ignored. '
             f'Data refreshed "&TEXT({a},"d mmm yyyy h:mm AM/PM")&" ET."')
    ready = f'{r.P("I")},"OPEN",{r.P("J")},FALSE,{r.P("D")},"<>Bot"'
    tiles = [
        (f'=COUNTIFS({ready})', "Open PRs (not drafts)",
         f'=COUNTIFS({r.P("I")},"OPEN",{r.P("J")},TRUE,{r.P("D")},"<>Bot")&" drafts besides · bots excluded"'),
        (f'=COUNTIFS({ready},{r.P("R")},"<"&({a}-30))', "PRs untouched 30+ days", '="by people · drafts not counted"'),
        (f'=COUNTIFS({ready},{r.P("P")},"APPROVED")', "PRs approved, not merged", '="ready to ship"'),
        (f'=COUNTIFS({ready},{r.P("P")},"CHANGES_REQUESTED")', "PRs with changes requested", '="waiting on authors"'),
        (f'=COUNTIFS({r.I("G")},"OPEN",{r.I("D")},"<>Bot")', "Open issues",
         f'=COUNTIFS({r.I("G")},"OPEN",{r.I("D")},"<>Bot",{r.I("I")},0)&" with no comments from people"'),
        (f'=COUNTIFS({r.I("G")},"OPEN",{r.I("D")},"<>Bot",{r.I("L")},"<"&({a}-30))', "Issues untouched 30+ days",
         '="by people"'),
        (f'=COUNTIFS({r.I("G")},"OPEN",{r.I("D")},"<>Bot",{r.I("N")},"")', "Open issues without a topic label",
         '="status/ labels not counted"'),
        (f'=COUNTIFS({r.I("G")},"OPEN",{r.I("D")},"External")', "Open issues from outside the team", '=""'),
    ]
    for k, (value, label, sub) in enumerate(tiles):
        tile(ws, 8 + 4 * (k // 4), 1 + 2 * (k % 4), value, label, sub)

    row = 17
    section(ws, row, "List", "Filtered by the yellow dropdowns. Numbers link to GitHub.")
    table_header(ws, row + 1, ["#", "Title", "Author", "Affiliation", "Untouched (days)", "", "Opened"])
    ws.cell(row=row + 1, column=6, value='=IF($A$6="Pull requests","Review state","Comments (people)")')
    min_days = 'SWITCH($E$6,"Any time",0,"7 days",7,"30 days",30,"90 days",90)'
    order = 'IF($I$6="Newest first",7,5)'

    def listing(cols: dict[str, str], keep: str) -> str:
        return (f'ARRAYFORMULA(LET(t,HSTACK({cols["link"]},{cols["title"]},{cols["author"]},'
                f'{cols["aff"]},INT({a}-{cols["touched"]}),{cols["extra"]},{cols["opened"]}),'
                f'SORT(FILTER(t,{keep}),{order},FALSE)))')

    pr = listing(
        {"link": f'HYPERLINK({r.P("Q")},"#"&{r.P("A")})', "title": r.P("B"), "author": r.P("C"),
         "aff": r.P("D"), "touched": r.P("R"), "opened": f'TEXT({r.P("F")},"yyyy-mm-dd")',
         "extra": f'IF({r.P("J")},"draft",IF({r.P("P")}="","no decision",LOWER(SUBSTITUTE({r.P("P")},"_"," "))))'},
        f'({r.P("I")}="OPEN")*((($G$6="Include drafts")+({r.P("J")}=FALSE))>0)*'
        f'(({a}-{r.P("R")})>={min_days})*{aff_ok(r.P("D"), "$C$6")}')
    iss = listing(
        {"link": f'HYPERLINK({r.I("K")},"#"&{r.I("A")})', "title": r.I("B"), "author": r.I("C"),
         "aff": r.I("D"), "touched": r.I("L"), "opened": f'TEXT({r.I("E")},"yyyy-mm-dd")', "extra": r.I("I")},
        f'({r.I("G")}="OPEN")*(({a}-{r.I("L")})>={min_days})*{aff_ok(r.I("D"), "$C$6")}')
    ws.cell(row=row + 2, column=1, value=f'=IFERROR(IF($A$6="Pull requests",{pr},{iss}),"Nothing matches these filters.")')
    for col, width in {"A": 9, "B": 52, "C": 18, "D": 17, "E": 13, "F": 17, "G": 13}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A7"

# ---------- Overview ----------

OVERVIEW_WEEKS = 26
OVERVIEW_MONTHS = 36


def span_label(start: str, end: str) -> str:
    """Array formula text: "14–21 Sep" for the days a bucket covers (end excluded)."""
    last = f"({end}-1/86400)"
    return (f'IF(MONTH({start})=MONTH({last}),TEXT({start},"d")&"–"&TEXT({last},"d mmm"),'
            f'TEXT({start},"d mmm")&"–"&TEXT({last},"d mmm"))')


def line_chart(ws, title: str, rows: tuple[int, int], series: list[tuple[int, str, str]], anchor: str,
               width: float = 12.5, height: float = 7.5, bar: bool = False, category_col: int = 1,
               x_title: str | None = None) -> None:
    chart = BarChart() if bar else LineChart()
    if bar:
        chart.type, chart.grouping = "col", "clustered"
    chart.title = title
    chart.height, chart.width = height, width
    chart.legend.position = "b"
    first, last = rows
    for col, name, color in series:
        s = Series(Reference(ws, min_col=col, min_row=first, max_row=last), title=name)
        s.graphicalProperties.line.solidFill = color
        if bar:
            s.graphicalProperties.solidFill = color
        else:
            s.graphicalProperties.line.width = 28575  # 2.25 pt
            s.smooth = False
        chart.series.append(s)
    chart.set_categories(Reference(ws, min_col=category_col, min_row=first, max_row=last))
    if x_title:
        chart.x_axis.title = x_title
    ws.add_chart(chart, anchor)


def overview(ws, r: Ranges) -> None:
    """A fixed view, the same for everyone: last week, the last 26 weeks, and every month."""
    a = r.asof
    ws["Z1"] = "helper"
    ws["Z2"] = f"=LET(d,INT({a})-WEEKDAY(INT({a}),3)+0.5,IF(d>{a},d-7,d))"  # end of the last full week
    ws["Z2"].number_format = "yyyy-mm-dd hh:mm"
    ws.column_dimensions["Z"].hidden = True
    end = "$Z$2"
    masthead(ws, '="Overview"',
             f'="Last full week: "&TEXT({end}-7,"d mmm")&" – "&TEXT({end},"d mmm yyyy")&", Monday noon to Monday '
             f'noon US Eastern. The same for everyone: this tab has no filters. Bots are excluded. Data refreshed "'
             f'&TEXT({a},"d mmm yyyy h:mm AM/PM")&" ET."')

    def week_count(rng: str, extra: str, weeks_back: int) -> str:
        s, e = f"({end}-{7 * (weeks_back + 1)})", f"({end}-{7 * weeks_back})"
        return f'COUNTIFS({rng},">="&{s},{rng},"<"&{e}{extra})'

    people_pr = f',{r.P("D")},"<>Bot"'
    measures = {
        "PRs opened": (r.P("F"), people_pr),
        "PRs merged": (r.P("G"), people_pr),
        "PRs reviewed": (r.R("E"), f',{r.R("C")},"<>Bot",{r.R("F")},TRUE'),
        "Issues opened": (r.I("E"), f',{r.I("D")},"<>Bot"'),
        "Issues closed": (r.I("F"), f',{r.I("D")},"<>Bot"'),
    }

    def period_tile(k: int, name: str) -> None:
        rng, extra = measures[name]
        last, prev = week_count(rng, extra, 0), week_count(rng, extra, 1)
        avg = "(" + "+".join(week_count(rng, extra, w) for w in range(4)) + ")/4"
        tile(ws, 5 + 4 * (k // 4), 1 + 2 * (k % 4), f"={last}", f"{name} last week",
             f'="previous week "&{prev}&" · 4-week average "&TEXT({avg},"0.0")')

    for k, name in enumerate(["PRs opened", "PRs merged", "PRs reviewed"]):
        period_tile(k, name)
    median = (f'IFERROR(MEDIAN(FILTER({r.P("M")},{r.P("F")}>={end}-7,{r.P("F")}<{end},{r.P("M")}<>"",'
              f'{r.P("D")}<>"Bot")),"–")')
    median_prev = median.replace(f"{end}-7", f"{end}-14").replace(f"<{end},", f"<{end}-7,")
    tile(ws, 5, 7, f"={median}", "Median hours to first review", f'="previous week "&IFERROR(TEXT({median_prev},"0.0"),"–")', "0.0")
    for k, name in enumerate(["Issues opened", "Issues closed"]):
        period_tile(4 + k, name)
    ready = f'{r.P("I")},"OPEN",{r.P("J")},FALSE,{r.P("D")},"<>Bot"'
    tile(ws, 9, 5, f"=COUNTIFS({ready})", "Open PRs now (not drafts)",
         f'=COUNTIFS({ready},{r.P("R")},"<"&({a}-30))&" untouched by people for 30+ days"')
    tile(ws, 9, 7, f'=COUNTIFS({r.I("G")},"OPEN",{r.I("D")},"<>Bot")', "Open issues now",
         f'=COUNTIFS({r.I("G")},"OPEN",{r.I("D")},"<>Bot",{r.I("L")},"<"&({a}-30))&" untouched by people for 30+ days"')

    # --- Weekly table (feeds the charts) ---
    wk_head, wk_first = 68, 69
    wk_last = wk_first + OVERVIEW_WEEKS - 1
    section(ws, wk_head - 1, f"Weekly figures, last {OVERVIEW_WEEKS} weeks",
            "Each row is Monday noon to Monday noon US Eastern. The charts above come from this table.")
    table_header(ws, wk_head, ["Week", "PRs opened", "PRs merged", "PRs closed without merging", "PRs reviewed",
                               "Issues opened", "Issues closed", "Open issues at week end", "Open PRs at week end",
                               "Opened by Digital Asset", "Opened by Canton Foundation", "Opened by External",
                               "Median hours to first review", "Week starting (chart label)"])
    P, R_, I = r.P, r.R, r.I

    def open_at(created: str, closed: str, who: str, at: str) -> str:
        return (f'COUNTIFS({created},"<"&{at},{closed},"",{who},"<>Bot")'
                f'+COUNTIFS({created},"<"&{at},{closed},">="&{at},{who},"<>Bot")')

    ws.cell(row=wk_first, column=1, value=(
        f'=ARRAYFORMULA(LET(k,SEQUENCE({OVERVIEW_WEEKS}),st,{end}-7*({OVERVIEW_WEEKS}+1-k),en,st+7,'
        f'HSTACK({span_label("st", "en")},'
        f'COUNTIFS({P("F")},">="&st,{P("F")},"<"&en,{P("D")},"<>Bot"),'
        f'COUNTIFS({P("G")},">="&st,{P("G")},"<"&en,{P("D")},"<>Bot"),'
        f'COUNTIFS({P("H")},">="&st,{P("H")},"<"&en,{P("G")},"",{P("D")},"<>Bot"),'
        f'COUNTIFS({R_("E")},">="&st,{R_("E")},"<"&en,{R_("C")},"<>Bot",{R_("F")},TRUE),'
        f'COUNTIFS({I("E")},">="&st,{I("E")},"<"&en,{I("D")},"<>Bot"),'
        f'COUNTIFS({I("F")},">="&st,{I("F")},"<"&en,{I("D")},"<>Bot"),'
        f'{open_at(I("E"), I("F"), I("D"), "en")},'
        f'{open_at(P("F"), P("H"), P("D"), "en")},'
        f'COUNTIFS({P("F")},">="&st,{P("F")},"<"&en,{P("D")},"Digital Asset"),'
        f'COUNTIFS({P("F")},">="&st,{P("F")},"<"&en,{P("D")},"Canton Foundation"),'
        f'COUNTIFS({P("F")},">="&st,{P("F")},"<"&en,{P("D")},"External"),'
        f'MAP(st,en,LAMBDA(a,b,IFERROR(MEDIAN(FILTER({P("M")},{P("F")}>=a,{P("F")}<b,{P("M")}<>"",'
        f'{P("D")}<>"Bot")),""))),'
        f'TEXT(st,"d mmm"))))'))
    for i in range(wk_first, wk_last + 1):
        ws.cell(row=i, column=13).number_format = "0.0"

    # --- Monthly table, from the first month with activity ---
    mo_head = wk_last + 4
    mo_first = mo_head + 1
    mo_last = mo_first + OVERVIEW_MONTHS - 1
    section(ws, mo_head - 1, "Monthly figures since the start", "Calendar months, US Eastern. The current month is to date.")
    table_header(ws, mo_head, ["Month", "PRs opened", "PRs merged", "PRs closed without merging", "PRs reviewed",
                               "Issues opened", "Issues closed", "Open issues at month end"])
    first_day = f"INT(MIN({P('F')},{I('E')}))"
    m0 = f"DATE(YEAR({first_day}),MONTH({first_day}),1)"
    n = f"MIN({OVERVIEW_MONTHS},(YEAR({a})-YEAR({first_day}))*12+MONTH({a})-MONTH({first_day})+1)"
    ws.cell(row=mo_first, column=1, value=(
        f'=ARRAYFORMULA(LET(k,SEQUENCE({n})-1,st,EDATE({m0},k),nx,EDATE({m0},k+1),en,IF(nx>{a},{a},nx),'
        f'HSTACK(TEXT(st,"mmm yyyy")&IF(nx>{a}," (to date)",""),'
        f'COUNTIFS({P("F")},">="&st,{P("F")},"<"&en,{P("D")},"<>Bot"),'
        f'COUNTIFS({P("G")},">="&st,{P("G")},"<"&en,{P("D")},"<>Bot"),'
        f'COUNTIFS({P("H")},">="&st,{P("H")},"<"&en,{P("G")},"",{P("D")},"<>Bot"),'
        f'COUNTIFS({R_("E")},">="&st,{R_("E")},"<"&en,{R_("C")},"<>Bot",{R_("F")},TRUE),'
        f'COUNTIFS({I("E")},">="&st,{I("E")},"<"&en,{I("D")},"<>Bot"),'
        f'COUNTIFS({I("F")},">="&st,{I("F")},"<"&en,{I("D")},"<>Bot"),'
        f'{open_at(I("E"), I("F"), I("D"), "en")})))'))

    # --- Charts ---
    blue, mid, light, slate = "1F4E8C", "4A86C8", "A9C6E8", "5B6C8F"
    orange, peach = "E07B39", "F2B880"
    wk = (wk_first, wk_last)
    weekly = {"category_col": 14, "x_title": "Week starting Monday (noon to noon ET)"}
    section(ws, 13, f"Weekly trends, last {OVERVIEW_WEEKS} weeks")
    line_chart(ws, "Pull requests per week", wk, [(2, "Opened", blue), (3, "Merged", mid), (5, "Reviewed", slate)],
               "A14", **weekly)
    line_chart(ws, "Issues per week", wk, [(6, "Opened", orange), (7, "Closed", peach)], "F14", **weekly)
    line_chart(ws, "Open at the end of each week", wk, [(9, "Open PRs", blue), (8, "Open issues", orange)], "A31",
               **weekly)
    line_chart(ws, "PRs opened per week, by affiliation", wk,
               [(10, "Digital Asset", blue), (11, "Canton Foundation", orange), (12, "External", "2E9E6B")], "F31",
               **weekly)
    section(ws, 48, "Since the start, by month")
    line_chart(ws, "PRs and issues per month", (mo_first, mo_last),
               [(2, "PRs opened", blue), (3, "PRs merged", mid), (6, "Issues opened", orange), (7, "Issues closed", peach)],
               "A49", width=25, height=8, bar=True)
    ws.column_dimensions["A"].width = 19
    ws.freeze_panes = "A4"


def add_reports(wb, r: Ranges) -> None:
    for title, builder in (("Overview", overview), ("Activity", activity), ("People", people),
                           ("Open work", open_work)):
        ws = wb[title]
        builder(ws, r)
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = "1F3A5F"
